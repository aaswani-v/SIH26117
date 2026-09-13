import json
import os
from datetime import datetime, timezone

from docx import Document


def write_json_log(result, out_dir: str = "./output") -> str:
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"run_{_timestamp()}.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(
            {
                "route": result.route,
                "model_used": result.model_used,
                "retrieved_sources": result.retrieved_sources,
                "steps_log": result.steps_log,
                "output": result.output,
            },
            f,
            indent=2,
        )
    return path


def write_docx(result, out_dir: str = "./output") -> str:
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"result_{_timestamp()}.docx")

    doc = Document()
    doc.add_heading("Sovereign AI Workbench — Task Result", level=1)
    doc.add_paragraph(f"Route: {result.route}  |  Model: {result.model_used}")
    if result.retrieved_sources:
        doc.add_heading("Sources consulted", level=2)
        for src in result.retrieved_sources:
            doc.add_paragraph(src, style="List Bullet")
    doc.add_heading("Result", level=2)
    doc.add_paragraph(result.output)
    doc.save(path)
    return path


def _timestamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
